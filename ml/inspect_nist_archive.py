#!/usr/bin/env python3
"""Safely inventory NIST MaTTO archives and inspect candidate tabular specimen data.
This script NEVER infers engineering labels or marks a dataset as training-ready.
"""
import argparse, csv, io, json, zipfile, hashlib, re
from pathlib import Path
MAX_MEMBER=40*1024*1024
def sniff_zip(path:Path):
 records=[]
 with zipfile.ZipFile(path) as z:
  for member in z.infolist():
   if member.is_dir():continue
   item={"name":member.filename,"bytes":member.file_size}
   if member.file_size>MAX_MEMBER:
    item["status"]="skipped_large";records.append(item);continue
   suffix=Path(member.filename).suffix.lower()
   if suffix in (".csv",".tsv",".txt"):
    b=z.read(member)
    try:
     text=b.decode("utf-8-sig")
    except UnicodeDecodeError:
     item["status"]="non_utf8";records.append(item);continue
    delimiter="\\t" if suffix==".tsv" else ","
    lines=[line for line in text.splitlines() if line.strip()]
    sample=lines[:4]
    item.update(status="preview_only",sample_lines=[x[:250] for x in sample],sha256=hashlib.sha256(b).hexdigest())
   elif suffix==".xlsx":
    try:
     import openpyxl
     b=z.read(member)
     wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
     sheets=[]
     for sh in wb:
      first=list(sh.iter_rows(min_row=1,max_row=4,values_only=True))
      sheets.append({"name":sh.title,"rows":sh.max_row,"columns":sh.max_column,"sample":[[str(v)[:100] if v is not None else None for v in row[:15]] for row in first]})
     item.update(status="preview_only",sheets=sheets,sha256=hashlib.sha256(b).hexdigest())
     wb.close()
    except ImportError:item["status"]="openpyxl_not_installed"
   else:item["status"]="other_format"
   records.append(item)
 return records
def main():
 p=argparse.ArgumentParser();p.add_argument("--zip",required=True,type=Path);p.add_argument("--out",default="ml/data/nist_audit/specimen_inventory.json",type=Path);a=p.parse_args()
 if not a.zip.is_file():raise SystemExit("Archive not found: "+str(a.zip))
 summary={"dataset":"NIST MaTTO experimental tensile measurements","archive_sha256":hashlib.sha256(a.zip.read_bytes()).hexdigest(),"training_ready":False,"reason":"Data fields and specimen orientation conventions require manual mapping and provenance review before training","members":sniff_zip(a.zip)}
 a.out.parent.mkdir(parents=True,exist_ok=True);a.out.write_text(json.dumps(summary,indent=2)+"\n")
 print(json.dumps({"output":str(a.out),"members":len(summary["members"]),"training_ready":False},indent=2))
if __name__=="__main__":main()
